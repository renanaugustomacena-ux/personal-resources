#!/usr/bin/env python3
"""
Advanced Presentation Engine for Security Domain Documentation.
Generates rich, dense PPTX presentations with sophisticated visuals.
"""

import io
import math
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch
import numpy as np

# ════════════════════════════════════════════════════════════════════
#  COLOR PALETTE — Cyberpunk / Dark Professional Theme
# ════════════════════════════════════════════════════════════════════
class Colors:
    BG_DEEP     = RGBColor(0x0A, 0x0A, 0x1A)
    BG_CARD     = RGBColor(0x14, 0x14, 0x2A)
    BG_ELEVATED = RGBColor(0x1C, 0x1C, 0x34)
    BG_HOVER    = RGBColor(0x24, 0x24, 0x40)
    BG_ACCENT   = RGBColor(0x1A, 0x1A, 0x2E)

    CYAN        = RGBColor(0x00, 0xD4, 0xFF)
    CYAN_DIM    = RGBColor(0x00, 0x8C, 0xB4)
    TEAL        = RGBColor(0x00, 0xE5, 0xA0)
    GREEN       = RGBColor(0x00, 0xFF, 0x88)
    GREEN_DIM   = RGBColor(0x00, 0xAA, 0x5C)
    RED         = RGBColor(0xFF, 0x44, 0x44)
    RED_DIM     = RGBColor(0xCC, 0x22, 0x22)
    ORANGE      = RGBColor(0xFF, 0x8C, 0x00)
    YELLOW      = RGBColor(0xFF, 0xD7, 0x00)
    YELLOW_DIM  = RGBColor(0xCC, 0xAA, 0x00)
    PURPLE      = RGBColor(0xBB, 0x86, 0xFC)
    PURPLE_DIM  = RGBColor(0x88, 0x60, 0xCC)
    MAGENTA     = RGBColor(0xFF, 0x00, 0x80)
    PINK        = RGBColor(0xFF, 0x6B, 0xB5)
    BLUE        = RGBColor(0x44, 0x88, 0xFF)
    BLUE_DIM    = RGBColor(0x33, 0x66, 0xCC)
    INDIGO      = RGBColor(0x66, 0x44, 0xFF)

    WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
    LGRAY       = RGBColor(0xCC, 0xCC, 0xCC)
    MGRAY       = RGBColor(0x99, 0x99, 0x99)
    DGRAY       = RGBColor(0x55, 0x55, 0x66)
    BORDER      = RGBColor(0x33, 0x33, 0x50)

    SEVERITY_CRITICAL = RGBColor(0xFF, 0x22, 0x22)
    SEVERITY_HIGH     = RGBColor(0xFF, 0x66, 0x00)
    SEVERITY_MEDIUM   = RGBColor(0xFF, 0xCC, 0x00)
    SEVERITY_LOW      = RGBColor(0x00, 0xCC, 0x66)
    SEVERITY_INFO     = RGBColor(0x44, 0x88, 0xFF)

    ACCENT_CYCLE = [CYAN, GREEN, YELLOW, RED, PURPLE, ORANGE, TEAL, MAGENTA, BLUE, PINK, INDIGO]

    _NAME_MAP = None

    @classmethod
    def resolve(cls, c):
        """Accept RGBColor or string name ('red','cyan',...) and return RGBColor."""
        if isinstance(c, RGBColor):
            return c
        if c is None:
            return None
        if cls._NAME_MAP is None:
            cls._NAME_MAP = {
                "cyan": cls.CYAN, "green": cls.GREEN, "red": cls.RED,
                "yellow": cls.YELLOW, "purple": cls.PURPLE, "orange": cls.ORANGE,
                "teal": cls.TEAL, "magenta": cls.MAGENTA, "blue": cls.BLUE,
                "pink": cls.PINK, "indigo": cls.INDIGO, "white": cls.WHITE,
                "lgray": cls.LGRAY, "mgray": cls.MGRAY, "dgray": cls.DGRAY,
                "border": cls.BORDER, "bg": cls.BG_DEEP, "card": cls.BG_CARD,
                "elevated": cls.BG_ELEVATED, "hover": cls.BG_HOVER,
                "accent": cls.BG_ACCENT, "cyan_dim": cls.CYAN_DIM,
                "green_dim": cls.GREEN_DIM, "red_dim": cls.RED_DIM,
                "yellow_dim": cls.YELLOW_DIM, "purple_dim": cls.PURPLE_DIM,
                "blue_dim": cls.BLUE_DIM,
            }
        return cls._NAME_MAP.get(str(c).lower().replace(" ", "_"), cls.CYAN)

    # Matplotlib hex equivalents
    HEX = {
        "bg": "#0A0A1A", "card": "#14142A", "elevated": "#1C1C34",
        "cyan": "#00D4FF", "green": "#00FF88", "red": "#FF4444",
        "yellow": "#FFD700", "purple": "#BB86FC", "orange": "#FF8C00",
        "teal": "#00E5A0", "magenta": "#FF0080", "blue": "#4488FF",
        "pink": "#FF6BB5", "indigo": "#6644FF", "white": "#FFFFFF",
        "lgray": "#CCCCCC", "mgray": "#999999", "dgray": "#555566",
        "border": "#333350", "grid": "#222240",
    }

C = Colors


# ════════════════════════════════════════════════════════════════════
#  MATPLOTLIB THEME
# ════════════════════════════════════════════════════════════════════
def mpl_theme():
    plt.rcParams.update({
        "figure.facecolor": C.HEX["bg"],
        "axes.facecolor": C.HEX["card"],
        "axes.edgecolor": C.HEX["border"],
        "axes.labelcolor": C.HEX["lgray"],
        "text.color": C.HEX["lgray"],
        "xtick.color": C.HEX["mgray"],
        "ytick.color": C.HEX["mgray"],
        "grid.color": C.HEX["grid"],
        "grid.alpha": 0.4,
        "font.size": 11,
        "font.family": "sans-serif",
        "axes.spines.top": False,
        "axes.spines.right": False,
    })


def chart_to_image(fig, dpi=180):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", pad_inches=0.2)
    buf.seek(0)
    plt.close(fig)
    return buf


# ════════════════════════════════════════════════════════════════════
#  PRESENTATION ENGINE
# ════════════════════════════════════════════════════════════════════
class PresentationEngine:
    W = Inches(13.333)
    H = Inches(7.5)

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = self.W
        self.prs.slide_height = self.H
        self.slide_count = 0

    def save(self, path):
        self.prs.save(path)
        print(f"  [+] Saved: {path} ({self.slide_count} slides)")

    # ── Core Slide Creation ──────────────────────────────────────
    def _blank(self):
        layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(layout)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = C.BG_DEEP
        self.slide_count += 1
        return slide

    # ── Text Primitives ──────────────────────────────────────────
    def _text(self, slide, left, top, width, height, text, size=14,
              color=C.LGRAY, bold=False, italic=False, align=PP_ALIGN.LEFT,
              font_name=None, anchor=MSO_ANCHOR.TOP):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        try:
            tf.vertical_anchor = anchor
        except Exception:
            pass
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = C.resolve(color)
        p.font.bold = bold
        p.font.italic = italic
        if font_name:
            p.font.name = font_name
        p.alignment = align
        return txBox

    def _multiline(self, slide, left, top, width, height, lines,
                   size=13, color=C.LGRAY, line_spacing=1.15, bullet=""):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        first = True
        for line in lines:
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            p.text = f"{bullet}{line}" if bullet else line
            p.font.size = Pt(size)
            p.font.color.rgb = C.resolve(color)
            p.space_after = Pt(size * 0.3)
            try:
                p.line_spacing = Pt(size * line_spacing)
            except Exception:
                pass
        return txBox

    def _rich_text(self, slide, left, top, width, height, segments):
        """segments: list of (text, size, color, bold, italic)"""
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        for i, seg in enumerate(segments):
            text, sz, clr, bld, ital = seg
            if i == 0:
                run = p.runs[0] if p.runs else p.add_run()
                run.text = text
            else:
                run = p.add_run()
                run.text = text
            run.font.size = Pt(sz)
            run.font.color.rgb = C.resolve(clr)
            run.font.bold = bld
            run.font.italic = ital
        return txBox

    # ── Shape Primitives ─────────────────────────────────────────
    def _box(self, slide, left, top, w, h, fill_color, text="",
             font_size=11, font_color=C.WHITE, border_color=None,
             border_width=1.5, align=PP_ALIGN.CENTER, bold=True,
             shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE):
        shape = slide.shapes.add_shape(shape_type, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C.resolve(fill_color)
        if border_color:
            shape.line.color.rgb = C.resolve(border_color)
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        try:
            tf.vertical_anchor = anchor
        except Exception:
            pass
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = C.resolve(font_color)
        p.font.bold = bold
        p.alignment = align
        tf.margin_left = Pt(6)
        tf.margin_right = Pt(6)
        tf.margin_top = Pt(4)
        tf.margin_bottom = Pt(4)
        return shape

    def _accent_box(self, slide, left, top, w, h, accent_color, text="",
                    font_size=11, font_color=C.LGRAY, bold=True):
        """Box with colored left accent bar."""
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = C.resolve(accent_color)
        bar.line.fill.background()
        shape = self._box(slide, left + Inches(0.08), top, w - Inches(0.08), h,
                         C.BG_ELEVATED, text, font_size, font_color,
                         border_color=C.BORDER, bold=bold, align=PP_ALIGN.LEFT)
        return shape

    def _glow_box(self, slide, left, top, w, h, glow_color, text="",
                  font_size=12, font_color=C.WHITE):
        """Box with glow-effect border."""
        outer = self._box(slide, left - Pt(2), top - Pt(2),
                         w + Pt(4), h + Pt(4), C.BG_DEEP, "",
                         border_color=glow_color, border_width=2.5)
        outer.fill.background()
        inner = self._box(slide, left, top, w, h, C.BG_ELEVATED, text,
                         font_size, font_color, border_color=glow_color, border_width=1)
        return inner

    def _pill(self, slide, left, top, w, h, fill_color, text="",
              font_size=10, font_color=C.WHITE):
        """Rounded pill/badge shape."""
        return self._box(slide, left, top, w, h, fill_color, text,
                        font_size, font_color, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)

    # ── Connectors and Arrows ────────────────────────────────────
    def _arrow_right(self, slide, x1, y, x2, color=C.CYAN, size=Inches(0.25)):
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1, y - size/2,
                                       x2 - x1, size)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C.resolve(color)
        shape.line.fill.background()
        return shape

    def _arrow_down(self, slide, x, y1, y2, color=C.CYAN, size=Inches(0.25)):
        shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x - size/2, y1,
                                       size, y2 - y1)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C.resolve(color)
        shape.line.fill.background()
        return shape

    def _line(self, slide, x1, y1, x2, y2, color=C.CYAN, width=1.5):
        line = slide.shapes.add_connector(1, x1, y1, x2, y2)
        line.line.color.rgb = C.resolve(color)
        line.line.width = Pt(width)
        return line

    def _chevron(self, slide, left, top, w, h, fill_color, text="",
                 font_size=10, font_color=C.WHITE):
        shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C.resolve(fill_color)
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = C.resolve(font_color)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        return shape

    # ── Decorative Elements ──────────────────────────────────────
    def _header_bar(self, slide, text, y=Inches(0.3)):
        self._box(slide, Inches(0), Inches(0), self.W, Inches(0.08),
                 C.CYAN, "", border_color=None)
        self._text(slide, Inches(0.8), y, Inches(11), Inches(0.7),
                  text, size=26, color=C.WHITE, bold=True)
        self._line(slide, Inches(0.8), y + Inches(0.65),
                  Inches(12.5), y + Inches(0.65), C.CYAN, 1.0)

    def _sub_header(self, slide, text, y=Inches(1.05)):
        self._text(slide, Inches(0.8), y, Inches(11), Inches(0.4),
                  text, size=15, color=C.CYAN, italic=True)

    def _footer(self, slide, left_text="", right_text=""):
        self._line(slide, Inches(0.5), Inches(7.0), Inches(12.8), Inches(7.0),
                  C.BORDER, 0.5)
        if left_text:
            self._text(slide, Inches(0.8), Inches(7.05), Inches(5), Inches(0.3),
                      left_text, size=8, color=C.DGRAY)
        if right_text:
            self._text(slide, Inches(8), Inches(7.05), Inches(4.5), Inches(0.3),
                      right_text, size=8, color=C.DGRAY, align=PP_ALIGN.RIGHT)

    def _slide_number(self, slide, num, total=None):
        txt = f"{num}" if not total else f"{num} / {total}"
        self._text(slide, Inches(12.2), Inches(7.05), Inches(1), Inches(0.3),
                  txt, size=8, color=C.DGRAY, align=PP_ALIGN.RIGHT)

    def _corner_accent(self, slide, color=C.CYAN):
        self._box(slide, Inches(0), Inches(0), Inches(0.08), Inches(1.2),
                 color, border_color=None)
        self._box(slide, Inches(0), Inches(0), Inches(1.2), Inches(0.08),
                 color, border_color=None)

    def _progress_bar(self, slide, progress, y=Inches(7.35), color=C.CYAN):
        total_w = Inches(13.333)
        self._box(slide, Inches(0), y, total_w, Inches(0.06), C.BG_CARD)
        self._box(slide, Inches(0), y, Emu(int(total_w * progress)), Inches(0.06), color)

    # ── Chart Embedding ──────────────────────────────────────────
    def _embed_chart(self, slide, fig, left, top, width, height=None):
        buf = chart_to_image(fig)
        return slide.shapes.add_picture(buf, left, top, width, height)

    # ════════════════════════════════════════════════════════════
    #  HIGH-LEVEL SLIDE TYPES
    # ════════════════════════════════════════════════════════════

    def title_slide(self, title, subtitle, domain="", chapter=""):
        slide = self._blank()
        self._box(slide, Inches(0), Inches(0), self.W, Inches(0.06), C.CYAN)
        self._box(slide, Inches(0), Inches(7.44), self.W, Inches(0.06), C.CYAN)
        self._corner_accent(slide, C.CYAN)
        self._box(slide, Inches(0), Inches(0), Inches(0.06), self.H, C.CYAN)
        diamond = self._box(slide, Inches(6.2), Inches(1.0), Inches(0.9), Inches(0.9),
                           C.BG_DEEP, "◆", font_size=32, font_color=C.CYAN,
                           shape_type=MSO_SHAPE.RECTANGLE)
        if domain:
            self._pill(slide, Inches(5.0), Inches(2.2), Inches(3.3), Inches(0.4),
                      C.BG_ELEVATED, domain, font_size=11, font_color=C.CYAN)
        self._text(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(2.0),
                  title, size=38, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
        self._line(slide, Inches(4), Inches(4.9), Inches(9.3), Inches(4.9), C.CYAN, 1.5)
        self._text(slide, Inches(1.5), Inches(5.1), Inches(10.3), Inches(1.0),
                  subtitle, size=18, color=C.LGRAY, align=PP_ALIGN.CENTER)
        if chapter:
            self._text(slide, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.5),
                      chapter, size=12, color=C.DGRAY, align=PP_ALIGN.CENTER)
        return slide

    def section_slide(self, title, section_num=None, description=""):
        slide = self._blank()
        self._box(slide, Inches(0), Inches(3.2), self.W, Inches(0.04), C.PURPLE)
        if section_num is not None:
            self._box(slide, Inches(5.8), Inches(1.5), Inches(1.7), Inches(1.7),
                     C.BG_DEEP, str(section_num), font_size=48, font_color=C.PURPLE,
                     border_color=C.PURPLE, border_width=2,
                     shape_type=MSO_SHAPE.RECTANGLE)
        self._text(slide, Inches(1.5), Inches(3.5), Inches(10.3), Inches(1.2),
                  title, size=34, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
        if description:
            self._text(slide, Inches(2), Inches(4.8), Inches(9.3), Inches(1.0),
                      description, size=16, color=C.LGRAY, align=PP_ALIGN.CENTER)
        self._line(slide, Inches(5), Inches(6.0), Inches(8.3), Inches(6.0), C.PURPLE, 1)
        return slide

    def agenda_slide(self, items, title="AGENDA"):
        slide = self._blank()
        self._header_bar(slide, title)
        mid = (len(items) + 1) // 2
        col1 = items[:mid]
        col2 = items[mid:]
        y = Inches(1.3)
        for i, item in enumerate(col1):
            clr = C.ACCENT_CYCLE[i % len(C.ACCENT_CYCLE)]
            num_box = self._box(slide, Inches(0.8), y, Inches(0.5), Inches(0.42),
                               C.BG_ELEVATED, str(i+1), font_size=12,
                               font_color=clr, border_color=clr, border_width=1)
            self._text(slide, Inches(1.45), y, Inches(5), Inches(0.42),
                      item, size=12, color=C.LGRAY)
            y += Inches(0.48)
        y = Inches(1.3)
        for i, item in enumerate(col2):
            idx = mid + i
            clr = C.ACCENT_CYCLE[idx % len(C.ACCENT_CYCLE)]
            self._box(slide, Inches(7.0), y, Inches(0.5), Inches(0.42),
                     C.BG_ELEVATED, str(idx+1), font_size=12,
                     font_color=clr, border_color=clr, border_width=1)
            self._text(slide, Inches(7.65), y, Inches(5), Inches(0.42),
                      item, size=12, color=C.LGRAY)
            y += Inches(0.48)
        self._line(slide, Inches(6.5), Inches(1.2), Inches(6.5), Inches(6.8), C.BORDER, 0.5)
        return slide

    def content_slide(self, heading, bullets, sub_heading=""):
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        y = Inches(1.5) if sub_heading else Inches(1.2)
        for line in bullets:
            self._accent_box(slide, Inches(0.8), y, Inches(11.5), Inches(0.38),
                            C.CYAN, f"  {line}", font_size=12, font_color=C.LGRAY, bold=False)
            y += Inches(0.44)
        return slide

    def content_slide_rich(self, heading, blocks, sub_heading=""):
        """blocks: list of (accent_color, title, description)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        y = Inches(1.5) if sub_heading else Inches(1.2)
        for accent, title, desc in blocks:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(0.8), y, Inches(0.06), Inches(0.65))
            bar.fill.solid()
            bar.fill.fore_color.rgb = C.resolve(accent)
            bar.line.fill.background()
            self._text(slide, Inches(1.0), y, Inches(3.5), Inches(0.3),
                      title, size=13, color=accent, bold=True)
            self._text(slide, Inches(1.0), y + Inches(0.28), Inches(11), Inches(0.4),
                      desc, size=11, color=C.LGRAY)
            y += Inches(0.72)
        return slide

    def two_column_slide(self, heading, left_title, left_items, right_title, right_items,
                         left_color=C.CYAN, right_color=C.GREEN):
        slide = self._blank()
        self._header_bar(slide, heading)
        self._box(slide, Inches(0.6), Inches(1.2), Inches(5.8), Inches(0.45),
                 C.BG_ELEVATED, left_title, font_size=14, font_color=left_color,
                 border_color=left_color, border_width=1)
        y = Inches(1.8)
        for item in left_items:
            self._text(slide, Inches(0.8), y, Inches(5.4), Inches(0.35),
                      f"▸ {item}", size=11, color=C.LGRAY)
            y += Inches(0.38)
        self._box(slide, Inches(6.8), Inches(1.2), Inches(5.8), Inches(0.45),
                 C.BG_ELEVATED, right_title, font_size=14, font_color=right_color,
                 border_color=right_color, border_width=1)
        y = Inches(1.8)
        for item in right_items:
            self._text(slide, Inches(7.0), y, Inches(5.4), Inches(0.35),
                      f"▸ {item}", size=11, color=C.LGRAY)
            y += Inches(0.38)
        self._line(slide, Inches(6.5), Inches(1.2), Inches(6.5), Inches(6.8), C.BORDER, 0.5)
        return slide

    def three_column_slide(self, heading, columns):
        """columns: list of (title, color, items)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        col_w = Inches(3.7)
        for ci, (title, color, items) in enumerate(columns[:3]):
            x = Inches(0.6) + ci * (col_w + Inches(0.3))
            self._box(slide, x, Inches(1.2), col_w, Inches(0.42),
                     C.BG_ELEVATED, title, font_size=13, font_color=color,
                     border_color=color, border_width=1)
            y = Inches(1.8)
            for item in items:
                self._text(slide, x + Inches(0.15), y, col_w - Inches(0.3), Inches(0.32),
                          f"▸ {item}", size=10, color=C.LGRAY)
                y += Inches(0.35)
        return slide

    def diagram_flow_slide(self, heading, nodes, sub_heading="", arrows=True,
                           node_width=Inches(2.0), y_pos=Inches(2.0)):
        """nodes: list of (label, color) — horizontal flow."""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        n = len(nodes)
        if n == 0:
            return slide
        total_w = Inches(12.0)
        gap = Inches(0.4)
        nw = min(node_width, (total_w - gap * (n - 1)) / n)
        spacing = nw + gap
        start_x = Inches(0.65)
        for i, (label, color) in enumerate(nodes):
            x = start_x + i * spacing
            self._glow_box(slide, x, y_pos, nw, Inches(0.75), color, label,
                          font_size=10, font_color=C.WHITE)
            if arrows and i < n - 1:
                ax = x + nw + Pt(4)
                self._arrow_right(slide, ax, y_pos + Inches(0.35),
                                 ax + gap - Pt(8), color)
        return slide

    def diagram_flow_vertical_slide(self, heading, nodes, x_pos=Inches(5.5)):
        """nodes: list of (label, color) — vertical flow."""
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.4)
        gap = Inches(0.35)
        nh = Inches(0.6)
        nw = Inches(3.0)
        for i, (label, color) in enumerate(nodes):
            self._glow_box(slide, x_pos, y, nw, nh, color, label,
                          font_size=11, font_color=C.WHITE)
            if i < len(nodes) - 1:
                self._arrow_down(slide, x_pos + nw / 2, y + nh + Pt(3),
                                y + nh + gap - Pt(3), color)
            y += nh + gap
        return slide

    def diagram_layers_slide(self, heading, layers, sub_heading=""):
        """layers: list of (label, color, description) — top to bottom stack."""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        n = len(layers)
        total_h = Inches(5.0)
        start_y = Inches(1.5)
        layer_h = min(Inches(0.7), total_h / n)
        gap = Inches(0.05)
        for i, (label, color, desc) in enumerate(layers):
            y = start_y + i * (layer_h + gap)
            w = Inches(5.0) - i * Inches(0.15)
            x = Inches(1.5) + i * Inches(0.075)
            self._box(slide, x, y, w, layer_h, C.BG_ELEVATED, label,
                     font_size=12, font_color=color, border_color=color,
                     border_width=1.5, bold=True)
            self._text(slide, x + w + Inches(0.3), y, Inches(6), layer_h,
                      desc, size=11, color=C.LGRAY)
        return slide

    def table_slide(self, heading, headers, rows, col_colors=None, sub_heading=""):
        """Rich table with alternating rows."""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        n_cols = len(headers)
        n_rows = len(rows)
        table_w = Inches(11.7)
        col_w = table_w / n_cols
        start_x = Inches(0.8)
        start_y = Inches(1.5) if sub_heading else Inches(1.2)
        row_h = Inches(0.38)
        hdr_h = Inches(0.42)
        if not col_colors:
            col_colors = [C.ACCENT_CYCLE[i % len(C.ACCENT_CYCLE)] for i in range(n_cols)]
        for ci, hdr in enumerate(headers):
            x = start_x + ci * col_w
            self._box(slide, x, start_y, col_w - Pt(2), hdr_h,
                     C.BG_ELEVATED, hdr, font_size=11, font_color=col_colors[ci],
                     border_color=col_colors[ci], border_width=1, bold=True)
        for ri, row in enumerate(rows):
            y = start_y + hdr_h + Pt(3) + ri * (row_h + Pt(2))
            bg = C.BG_CARD if ri % 2 == 0 else C.BG_DEEP
            for ci, cell in enumerate(row):
                x = start_x + ci * col_w
                self._box(slide, x, y, col_w - Pt(2), row_h, bg, str(cell),
                         font_size=10, font_color=C.LGRAY, bold=False,
                         align=PP_ALIGN.LEFT)
        return slide

    def comparison_slide(self, heading, left_title, left_items, right_title, right_items,
                         left_color=C.RED, right_color=C.GREEN, vs_text="VS"):
        slide = self._blank()
        self._header_bar(slide, heading)
        self._box(slide, Inches(0.6), Inches(1.2), Inches(5.5), Inches(0.5),
                 C.BG_ELEVATED, left_title, font_size=16, font_color=left_color,
                 border_color=left_color, border_width=2)
        y = Inches(1.9)
        for item in left_items:
            self._accent_box(slide, Inches(0.6), y, Inches(5.5), Inches(0.36),
                            left_color, f"  {item}", font_size=11,
                            font_color=C.LGRAY, bold=False)
            y += Inches(0.42)
        self._box(slide, Inches(6.15), Inches(3.0), Inches(1.0), Inches(1.0),
                 C.BG_DEEP, vs_text, font_size=18, font_color=C.YELLOW,
                 border_color=C.YELLOW, border_width=2,
                 shape_type=MSO_SHAPE.RECTANGLE)
        self._box(slide, Inches(7.2), Inches(1.2), Inches(5.5), Inches(0.5),
                 C.BG_ELEVATED, right_title, font_size=16, font_color=right_color,
                 border_color=right_color, border_width=2)
        y = Inches(1.9)
        for item in right_items:
            self._accent_box(slide, Inches(7.2), y, Inches(5.5), Inches(0.36),
                            right_color, f"  {item}", font_size=11,
                            font_color=C.LGRAY, bold=False)
            y += Inches(0.42)
        return slide

    def code_slide(self, heading, code_text, language="", notes=None):
        slide = self._blank()
        self._header_bar(slide, heading)
        if language:
            self._pill(slide, Inches(10.5), Inches(0.35), Inches(2), Inches(0.3),
                      C.BG_ELEVATED, language, font_size=9, font_color=C.CYAN)
        code_bg = self._box(slide, Inches(0.6), Inches(1.2), Inches(11.8), Inches(4.8),
                           C.BG_CARD, "", border_color=C.BORDER, border_width=1)
        self._line(slide, Inches(0.6), Inches(1.55), Inches(12.4), Inches(1.55), C.BORDER, 0.5)
        dots_y = Inches(1.32)
        self._box(slide, Inches(0.85), dots_y, Inches(0.15), Inches(0.15),
                 C.RED, "", shape_type=MSO_SHAPE.OVAL)
        self._box(slide, Inches(1.1), dots_y, Inches(0.15), Inches(0.15),
                 C.YELLOW, "", shape_type=MSO_SHAPE.OVAL)
        self._box(slide, Inches(1.35), dots_y, Inches(0.15), Inches(0.15),
                 C.GREEN, "", shape_type=MSO_SHAPE.OVAL)
        self._text(slide, Inches(0.85), Inches(1.65), Inches(11.3), Inches(4.2),
                  code_text, size=10, color=C.LGRAY, font_name="Consolas")
        if notes:
            y = Inches(6.15)
            for note in notes[:3]:
                self._text(slide, Inches(0.8), y, Inches(11.5), Inches(0.3),
                          f"→ {note}", size=10, color=C.MGRAY, italic=True)
                y += Inches(0.28)
        return slide

    def stats_slide(self, heading, stats, sub_heading=""):
        """stats: list of (value, label, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        n = len(stats)
        cols = min(n, 4)
        rows_count = math.ceil(n / cols)
        card_w = Inches(2.8)
        card_h = Inches(1.6)
        gap_x = Inches(0.3)
        gap_y = Inches(0.3)
        total_w = cols * card_w + (cols - 1) * gap_x
        start_x = (self.W - total_w) / 2
        start_y = Inches(1.8) if sub_heading else Inches(1.5)
        for i, (value, label, color) in enumerate(stats):
            row = i // cols
            col = i % cols
            x = start_x + col * (card_w + gap_x)
            y = start_y + row * (card_h + gap_y)
            self._box(slide, x, y, card_w, card_h, C.BG_ELEVATED, "",
                     border_color=color, border_width=1.5)
            self._text(slide, x, y + Inches(0.2), card_w, Inches(0.7),
                      str(value), size=32, color=color, bold=True, align=PP_ALIGN.CENTER)
            self._text(slide, x, y + Inches(0.95), card_w, Inches(0.5),
                      label, size=11, color=C.LGRAY, align=PP_ALIGN.CENTER)
        return slide

    def timeline_slide(self, heading, events):
        """events: list of (year/label, description, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        n = len(events)
        if n == 0:
            return slide
        line_y = Inches(3.5)
        self._line(slide, Inches(0.5), line_y, Inches(12.8), line_y, C.BORDER, 2)
        spacing = Inches(12.0) / max(n - 1, 1)
        for i, (label, desc, color) in enumerate(events):
            x = Inches(0.8) + i * spacing
            dot = self._box(slide, x - Inches(0.1), line_y - Inches(0.1),
                           Inches(0.2), Inches(0.2), color, "",
                           shape_type=MSO_SHAPE.OVAL)
            above = (i % 2 == 0)
            if above:
                self._text(slide, x - Inches(0.5), line_y - Inches(1.2),
                          Inches(1.5), Inches(0.3), str(label), size=11,
                          color=color, bold=True, align=PP_ALIGN.CENTER)
                self._text(slide, x - Inches(0.7), line_y - Inches(0.9),
                          Inches(1.8), Inches(0.6), desc, size=9,
                          color=C.LGRAY, align=PP_ALIGN.CENTER)
                self._line(slide, x, line_y - Inches(0.35), x, line_y - Inches(0.1), color, 1)
            else:
                self._text(slide, x - Inches(0.5), line_y + Inches(0.4),
                          Inches(1.5), Inches(0.3), str(label), size=11,
                          color=color, bold=True, align=PP_ALIGN.CENTER)
                self._text(slide, x - Inches(0.7), line_y + Inches(0.7),
                          Inches(1.8), Inches(0.6), desc, size=9,
                          color=C.LGRAY, align=PP_ALIGN.CENTER)
                self._line(slide, x, line_y + Inches(0.1), x, line_y + Inches(0.35), color, 1)
        return slide

    def process_slide(self, heading, steps, sub_heading=""):
        """Chevron-style process flow. steps: list of (label, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        n = len(steps)
        chev_w = Inches(1.8)
        gap = Inches(0.15)
        total = n * chev_w + (n - 1) * gap
        start_x = max(Inches(0.3), (self.W - total) / 2)
        y = Inches(2.2)
        for i, (label, color) in enumerate(steps):
            x = start_x + i * (chev_w + gap)
            self._chevron(slide, x, y, chev_w, Inches(0.7), color, label, 10)
        return slide

    def risk_matrix_slide(self, heading, items=None):
        """5x5 risk matrix. items: list of (likelihood_idx, impact_idx, label)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        cell_w = Inches(1.8)
        cell_h = Inches(0.9)
        start_x = Inches(2.5)
        start_y = Inches(1.5)
        impact_labels = ["Negligible", "Minor", "Moderate", "Major", "Critical"]
        likelihood_labels = ["Rare", "Unlikely", "Possible", "Likely", "Almost Certain"]
        risk_colors = [
            [C.GREEN, C.GREEN, C.YELLOW_DIM, C.YELLOW, C.ORANGE],
            [C.GREEN, C.YELLOW_DIM, C.YELLOW, C.ORANGE, C.RED_DIM],
            [C.YELLOW_DIM, C.YELLOW, C.ORANGE, C.RED_DIM, C.RED],
            [C.YELLOW, C.ORANGE, C.RED_DIM, C.RED, C.SEVERITY_CRITICAL],
            [C.ORANGE, C.RED_DIM, C.RED, C.SEVERITY_CRITICAL, C.SEVERITY_CRITICAL],
        ]
        self._text(slide, start_x, start_y - Inches(0.4), cell_w * 5, Inches(0.3),
                  "IMPACT →", size=12, color=C.CYAN, bold=True, align=PP_ALIGN.CENTER)
        self._text(slide, start_x - Inches(1.8), start_y + Inches(1.5),
                  Inches(1.5), Inches(0.3),
                  "LIKELIHOOD →", size=12, color=C.CYAN, bold=True)
        for ci, imp in enumerate(impact_labels):
            self._text(slide, start_x + ci * cell_w, start_y - Inches(0.15),
                      cell_w, Inches(0.2), imp, size=9, color=C.LGRAY,
                      align=PP_ALIGN.CENTER)
        for ri, lik in enumerate(likelihood_labels):
            y = start_y + Inches(0.1) + (4 - ri) * cell_h
            self._text(slide, start_x - Inches(1.8), y, Inches(1.6), cell_h,
                      lik, size=9, color=C.LGRAY, align=PP_ALIGN.RIGHT)
            for ci in range(5):
                x = start_x + ci * cell_w
                rc = risk_colors[ri][ci]
                self._box(slide, x, y, cell_w - Pt(2), cell_h - Pt(2),
                         rc, "", border_color=C.BG_DEEP, border_width=1,
                         font_size=8)
        if items:
            for li, ii, label in items:
                x = start_x + ii * cell_w + Inches(0.2)
                y = start_y + Inches(0.1) + (4 - li) * cell_h + Inches(0.15)
                self._pill(slide, x, y, cell_w - Inches(0.4), Inches(0.4),
                          C.BG_DEEP, label, font_size=8, font_color=C.WHITE)
        return slide

    def takeaway_slide(self, items, heading="Key Takeaways"):
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.3)
        for i, item in enumerate(items):
            clr = C.ACCENT_CYCLE[i % len(C.ACCENT_CYCLE)]
            num = self._box(slide, Inches(0.8), y, Inches(0.45), Inches(0.45),
                           C.BG_DEEP, str(i+1), font_size=14, font_color=clr,
                           border_color=clr, border_width=1.5,
                           shape_type=MSO_SHAPE.RECTANGLE)
            self._accent_box(slide, Inches(1.4), y, Inches(11), Inches(0.45),
                            clr, f"  {item}", font_size=12, font_color=C.LGRAY, bold=False)
            y += Inches(0.55)
        return slide

    def quote_slide(self, text, attribution="", color=C.CYAN):
        slide = self._blank()
        self._box(slide, Inches(2), Inches(2), Inches(9.3), Inches(3.5),
                 C.BG_ELEVATED, "", border_color=color, border_width=1)
        self._text(slide, Inches(2.3), Inches(1.7), Inches(1), Inches(0.8),
                  "❝", size=48, color=color, bold=True)
        self._text(slide, Inches(2.5), Inches(2.5), Inches(8.5), Inches(2.0),
                  text, size=18, color=C.WHITE, italic=True, align=PP_ALIGN.CENTER)
        if attribution:
            self._text(slide, Inches(2.5), Inches(4.5), Inches(8.5), Inches(0.5),
                      f"— {attribution}", size=14, color=C.MGRAY, align=PP_ALIGN.CENTER)
        return slide

    def definition_slide(self, heading, definitions):
        """definitions: list of (term, definition, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.3)
        for term, defn, color in definitions:
            self._box(slide, Inches(0.8), y, Inches(2.5), Inches(0.55),
                     C.BG_ELEVATED, term, font_size=12, font_color=color,
                     border_color=color, border_width=1)
            self._text(slide, Inches(3.5), y + Inches(0.05), Inches(9), Inches(0.5),
                      defn, size=11, color=C.LGRAY)
            y += Inches(0.62)
        return slide

    def warning_slide(self, heading, warnings):
        """warnings: list of (severity, text) where severity is critical/high/medium/low"""
        slide = self._blank()
        self._header_bar(slide, heading)
        sev_colors = {
            "critical": C.SEVERITY_CRITICAL, "high": C.SEVERITY_HIGH,
            "medium": C.SEVERITY_MEDIUM, "low": C.SEVERITY_LOW,
            "info": C.SEVERITY_INFO,
        }
        y = Inches(1.3)
        for severity, text in warnings:
            clr = sev_colors.get(severity.lower(), C.MGRAY)
            self._pill(slide, Inches(0.8), y + Inches(0.05), Inches(1.2), Inches(0.32),
                      clr, severity.upper(), font_size=9, font_color=C.WHITE)
            self._text(slide, Inches(2.2), y, Inches(10.2), Inches(0.42),
                      text, size=12, color=C.LGRAY)
            y += Inches(0.5)
        return slide

    def icon_grid_slide(self, heading, items, cols=4):
        """items: list of (icon_char, label, description, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        n = len(items)
        rows_count = math.ceil(n / cols)
        card_w = Inches(2.8)
        card_h = Inches(2.2)
        gap = Inches(0.25)
        total_w = cols * card_w + (cols - 1) * gap
        start_x = (self.W - total_w) / 2
        start_y = Inches(1.4)
        for i, (icon, label, desc, color) in enumerate(items):
            row = i // cols
            col = i % cols
            x = start_x + col * (card_w + gap)
            y = start_y + row * (card_h + gap)
            self._box(slide, x, y, card_w, card_h, C.BG_ELEVATED, "",
                     border_color=color, border_width=1)
            self._text(slide, x, y + Inches(0.15), card_w, Inches(0.5),
                      icon, size=28, color=color, align=PP_ALIGN.CENTER)
            self._text(slide, x + Inches(0.15), y + Inches(0.7), card_w - Inches(0.3),
                      Inches(0.35), label, size=12, color=C.WHITE, bold=True,
                      align=PP_ALIGN.CENTER)
            self._text(slide, x + Inches(0.15), y + Inches(1.1), card_w - Inches(0.3),
                      Inches(0.9), desc, size=9, color=C.LGRAY, align=PP_ALIGN.CENTER)
        return slide

    def architecture_slide(self, heading, layers, annotations=None):
        """Full-width layered architecture diagram.
        layers: list of (label, color, width_fraction)
        annotations: list of (text, x_inches, y_inches)
        """
        slide = self._blank()
        self._header_bar(slide, heading)
        n = len(layers)
        total_h = Inches(5.2)
        start_y = Inches(1.4)
        layer_h = total_h / (n + 0.5)
        max_w = Inches(10.0)
        for i, (label, color, w_frac) in enumerate(layers):
            w = max_w * w_frac
            x = (self.W - w) / 2
            y = start_y + i * layer_h
            self._box(slide, x, y, w, layer_h - Pt(3), C.BG_ELEVATED, label,
                     font_size=13, font_color=color, border_color=color,
                     border_width=1.5, bold=True)
        if annotations:
            for text, ax, ay in annotations:
                self._text(slide, Inches(ax), Inches(ay), Inches(3), Inches(0.3),
                          text, size=9, color=C.MGRAY, italic=True)
        return slide

    def matrix_comparison_slide(self, heading, row_labels, col_labels, data, colors=None):
        """data: 2D list of strings. colors: optional 2D list of RGBColor."""
        slide = self._blank()
        self._header_bar(slide, heading)
        n_rows = len(row_labels)
        n_cols = len(col_labels)
        cell_w = Inches(10.5) / (n_cols + 1)
        cell_h = Inches(0.45)
        start_x = Inches(0.8)
        start_y = Inches(1.3)
        for ci, cl in enumerate(col_labels):
            x = start_x + (ci + 1) * cell_w
            self._box(slide, x, start_y, cell_w - Pt(2), cell_h,
                     C.BG_ELEVATED, cl, font_size=10, font_color=C.CYAN, bold=True)
        for ri, rl in enumerate(row_labels):
            y = start_y + (ri + 1) * (cell_h + Pt(2))
            self._box(slide, start_x, y, cell_w - Pt(2), cell_h,
                     C.BG_ELEVATED, rl, font_size=10, font_color=C.LGRAY,
                     bold=True, align=PP_ALIGN.LEFT)
            for ci in range(n_cols):
                x = start_x + (ci + 1) * cell_w
                val = data[ri][ci] if ri < len(data) and ci < len(data[ri]) else ""
                clr = colors[ri][ci] if colors and ri < len(colors) and ci < len(colors[ri]) else C.LGRAY
                bg = C.BG_CARD if ri % 2 == 0 else C.BG_DEEP
                self._box(slide, x, y, cell_w - Pt(2), cell_h, bg, str(val),
                         font_size=9, font_color=clr, bold=False)
        return slide

    # ════════════════════════════════════════════════════════════
    #  CHART SLIDES (matplotlib-backed)
    # ════════════════════════════════════════════════════════════

    def bar_chart_slide(self, heading, labels, values, colors_hex=None,
                        xlabel="", ylabel="", horizontal=True, sub_heading=""):
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        if not colors_hex:
            colors_hex = [C.HEX["cyan"]] * len(labels)
        if horizontal:
            bars = ax.barh(labels, values, color=colors_hex, edgecolor=C.HEX["border"], height=0.6)
            ax.set_xlabel(xlabel, fontsize=11)
            for bar, val in zip(bars, values):
                ax.text(bar.get_width() + max(values)*0.02, bar.get_y() + bar.get_height()/2,
                       f"{val}", va="center", fontsize=9, color=C.HEX["lgray"])
        else:
            bars = ax.bar(labels, values, color=colors_hex, edgecolor=C.HEX["border"], width=0.6)
            ax.set_ylabel(ylabel, fontsize=11)
            for bar, val in zip(bars, values):
                ax.text(bar.get_x() + bar.get_width()/2, bar.get_height() + max(values)*0.02,
                       f"{val}", ha="center", fontsize=9, color=C.HEX["lgray"])
        ax.grid(axis="x" if horizontal else "y", alpha=0.3)
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.4) if sub_heading else Inches(1.2),
                         Inches(10.8), Inches(5.5))
        return slide

    def radar_chart_slide(self, heading, categories, datasets, sub_heading=""):
        """datasets: list of (label, values, color_hex)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        N = len(categories)
        angles = np.linspace(0, 2 * np.pi, N, endpoint=False).tolist()
        angles += angles[:1]
        fig, ax = plt.subplots(figsize=(6, 6), subplot_kw=dict(polar=True))
        ax.set_facecolor(C.HEX["card"])
        fig.patch.set_facecolor(C.HEX["bg"])
        for label, values, color in datasets:
            vals = values + values[:1]
            ax.plot(angles, vals, 'o-', linewidth=2, label=label, color=color)
            ax.fill(angles, vals, alpha=0.15, color=color)
        ax.set_xticks(angles[:-1])
        ax.set_xticklabels(categories, size=9, color=C.HEX["lgray"])
        ax.set_yticklabels([])
        ax.spines['polar'].set_color(C.HEX["border"])
        ax.grid(color=C.HEX["grid"], alpha=0.4)
        ax.legend(loc='upper right', bbox_to_anchor=(1.3, 1.1), fontsize=9,
                 facecolor=C.HEX["card"], edgecolor=C.HEX["border"],
                 labelcolor=C.HEX["lgray"])
        self._embed_chart(slide, fig, Inches(3.5), Inches(1.2), Inches(6.3), Inches(6.0))
        return slide

    def pie_chart_slide(self, heading, labels, values, colors_hex=None, sub_heading=""):
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(7, 5))
        if not colors_hex:
            colors_hex = [C.HEX["cyan"], C.HEX["green"], C.HEX["yellow"],
                         C.HEX["red"], C.HEX["purple"], C.HEX["orange"],
                         C.HEX["teal"], C.HEX["magenta"]][:len(labels)]
        wedges, texts, autotexts = ax.pie(values, labels=labels, colors=colors_hex,
                                          autopct='%1.1f%%', pctdistance=0.8,
                                          textprops={'fontsize': 10, 'color': C.HEX["lgray"]})
        for at in autotexts:
            at.set_color(C.HEX["white"])
            at.set_fontsize(9)
        self._embed_chart(slide, fig, Inches(3), Inches(1.2), Inches(7), Inches(5.5))
        return slide

    def heatmap_slide(self, heading, x_labels, y_labels, data, sub_heading=""):
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 5))
        data_arr = np.array(data)
        im = ax.imshow(data_arr, cmap='RdYlGn_r', aspect='auto')
        ax.set_xticks(range(len(x_labels)))
        ax.set_xticklabels(x_labels, fontsize=9, rotation=45, ha='right')
        ax.set_yticks(range(len(y_labels)))
        ax.set_yticklabels(y_labels, fontsize=9)
        for i in range(len(y_labels)):
            for j in range(len(x_labels)):
                ax.text(j, i, f"{data_arr[i, j]:.0f}", ha="center", va="center",
                       fontsize=8, color=C.HEX["white"])
        cbar = fig.colorbar(im, ax=ax, fraction=0.03, pad=0.04)
        cbar.ax.tick_params(labelsize=8, colors=C.HEX["lgray"])
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.2), Inches(10.8), Inches(5.5))
        return slide

    def multi_bar_chart_slide(self, heading, categories, datasets, ylabel="", sub_heading=""):
        """datasets: list of (label, values, color_hex)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        n = len(datasets)
        x = np.arange(len(categories))
        width = 0.8 / n
        for i, (label, values, color) in enumerate(datasets):
            offset = (i - n/2 + 0.5) * width
            ax.bar(x + offset, values, width, label=label, color=color,
                  edgecolor=C.HEX["border"])
        ax.set_xticks(x)
        ax.set_xticklabels(categories, fontsize=9)
        ax.set_ylabel(ylabel, fontsize=11)
        ax.legend(fontsize=9, facecolor=C.HEX["card"], edgecolor=C.HEX["border"],
                 labelcolor=C.HEX["lgray"])
        ax.grid(axis="y", alpha=0.3)
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.2), Inches(10.8), Inches(5.5))
        return slide

    def line_chart_slide(self, heading, x_data, datasets, xlabel="", ylabel="", sub_heading=""):
        """datasets: list of (label, y_values, color_hex)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        for label, y_vals, color in datasets:
            ax.plot(x_data, y_vals, 'o-', label=label, color=color, linewidth=2, markersize=5)
        ax.set_xlabel(xlabel, fontsize=11)
        ax.set_ylabel(ylabel, fontsize=11)
        ax.legend(fontsize=9, facecolor=C.HEX["card"], edgecolor=C.HEX["border"],
                 labelcolor=C.HEX["lgray"])
        ax.grid(True, alpha=0.3)
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.2), Inches(10.8), Inches(5.5))
        return slide

    def scatter_chart_slide(self, heading, datasets, xlabel="", ylabel="", sub_heading=""):
        """datasets: list of (label, x_vals, y_vals, color_hex, sizes)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        for label, xv, yv, color, sizes in datasets:
            ax.scatter(xv, yv, c=color, s=sizes, label=label, alpha=0.7, edgecolors=C.HEX["border"])
        ax.set_xlabel(xlabel, fontsize=11)
        ax.set_ylabel(ylabel, fontsize=11)
        ax.legend(fontsize=9, facecolor=C.HEX["card"], edgecolor=C.HEX["border"],
                 labelcolor=C.HEX["lgray"])
        ax.grid(True, alpha=0.3)
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.2), Inches(10.8), Inches(5.5))
        return slide

    def stacked_bar_slide(self, heading, categories, datasets, ylabel="", sub_heading=""):
        """datasets: list of (label, values, color_hex) — stacked."""
        slide = self._blank()
        self._header_bar(slide, heading)
        if sub_heading:
            self._sub_header(slide, sub_heading)
        mpl_theme()
        fig, ax = plt.subplots(figsize=(10, 4.5))
        x = np.arange(len(categories))
        bottom = np.zeros(len(categories))
        for label, values, color in datasets:
            ax.bar(x, values, 0.6, bottom=bottom, label=label, color=color,
                  edgecolor=C.HEX["border"])
            bottom += np.array(values)
        ax.set_xticks(x)
        ax.set_xticklabels(categories, fontsize=9)
        ax.set_ylabel(ylabel, fontsize=11)
        ax.legend(fontsize=9, facecolor=C.HEX["card"], edgecolor=C.HEX["border"],
                 labelcolor=C.HEX["lgray"])
        ax.grid(axis="y", alpha=0.3)
        self._embed_chart(slide, fig, Inches(1.2), Inches(1.2), Inches(10.8), Inches(5.5))
        return slide

    # ════════════════════════════════════════════════════════════
    #  COMPOUND SLIDE TYPES (content + visual)
    # ════════════════════════════════════════════════════════════

    def split_content_diagram(self, heading, bullets, nodes, bullet_width=Inches(5.5)):
        """Left side: bullet content, Right side: vertical flow diagram."""
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.3)
        for line in bullets:
            self._text(slide, Inches(0.8), y, bullet_width, Inches(0.35),
                      f"▸ {line}", size=11, color=C.LGRAY)
            y += Inches(0.38)
        diag_x = bullet_width + Inches(1.5)
        diag_w = Inches(2.5)
        y = Inches(1.3)
        gap = Inches(0.3)
        nh = Inches(0.55)
        for i, (label, color) in enumerate(nodes):
            self._box(slide, diag_x, y, diag_w, nh, C.BG_ELEVATED, label,
                     font_size=10, font_color=color, border_color=color, border_width=1)
            if i < len(nodes) - 1:
                self._arrow_down(slide, diag_x + diag_w / 2, y + nh + Pt(2),
                                y + nh + gap - Pt(2), color)
            y += nh + gap
        return slide

    def split_content_chart(self, heading, bullets, fig, chart_left=True):
        """Half bullets, half embedded chart."""
        slide = self._blank()
        self._header_bar(slide, heading)
        if chart_left:
            self._embed_chart(slide, fig, Inches(0.5), Inches(1.3), Inches(6), Inches(5.5))
            y = Inches(1.3)
            for line in bullets:
                self._text(slide, Inches(7), y, Inches(5.5), Inches(0.35),
                          f"▸ {line}", size=11, color=C.LGRAY)
                y += Inches(0.38)
        else:
            y = Inches(1.3)
            for line in bullets:
                self._text(slide, Inches(0.8), y, Inches(5.5), Inches(0.35),
                          f"▸ {line}", size=11, color=C.LGRAY)
                y += Inches(0.38)
            self._embed_chart(slide, fig, Inches(6.8), Inches(1.3), Inches(6), Inches(5.5))
        return slide

    def detail_slide(self, heading, blocks):
        """Dense detail slide. blocks: list of (title, content_lines, color)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.2)
        for title, lines, color in blocks:
            self._text(slide, Inches(0.8), y, Inches(4), Inches(0.3),
                      title, size=14, color=color, bold=True)
            y += Inches(0.35)
            for line in lines:
                self._text(slide, Inches(1.0), y, Inches(11.3), Inches(0.3),
                          f"  {line}", size=11, color=C.LGRAY)
                y += Inches(0.32)
            y += Inches(0.15)
        return slide

    def cross_reference_slide(self, items, heading="Cross-Reference Map"):
        """items: list of (domain_ref, description)"""
        slide = self._blank()
        self._header_bar(slide, heading)
        y = Inches(1.3)
        for i, (ref, desc) in enumerate(items):
            clr = C.ACCENT_CYCLE[i % len(C.ACCENT_CYCLE)]
            self._pill(slide, Inches(0.8), y + Inches(0.04), Inches(1.8), Inches(0.32),
                      C.BG_ELEVATED, ref, font_size=9, font_color=clr)
            self._text(slide, Inches(2.8), y, Inches(9.5), Inches(0.38),
                      desc, size=11, color=C.LGRAY)
            y += Inches(0.44)
        return slide

    def summary_slide(self, heading, summary_text, key_points=None):
        slide = self._blank()
        self._header_bar(slide, heading)
        self._text(slide, Inches(0.8), Inches(1.3), Inches(11.5), Inches(2.5),
                  summary_text, size=14, color=C.LGRAY)
        if key_points:
            y = Inches(4.0)
            self._text(slide, Inches(0.8), y - Inches(0.4), Inches(3), Inches(0.3),
                      "Key Points:", size=14, color=C.WHITE, bold=True)
            for i, point in enumerate(key_points):
                clr = C.ACCENT_CYCLE[i % len(C.ACCENT_CYCLE)]
                self._accent_box(slide, Inches(0.8), y, Inches(11.5), Inches(0.38),
                                clr, f"  {point}", font_size=11,
                                font_color=C.LGRAY, bold=False)
                y += Inches(0.44)
        return slide

    def end_slide(self, title="End of Presentation", subtitle=""):
        slide = self._blank()
        self._box(slide, Inches(0), Inches(0), self.W, Inches(0.06), C.CYAN)
        self._box(slide, Inches(0), Inches(7.44), self.W, Inches(0.06), C.CYAN)
        self._text(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(1.5),
                  title, size=36, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
        if subtitle:
            self._text(slide, Inches(1.5), Inches(4.5), Inches(10.3), Inches(1.0),
                      subtitle, size=16, color=C.LGRAY, align=PP_ALIGN.CENTER)
        return slide
